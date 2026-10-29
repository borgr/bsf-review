"""Build research_plan.docx from research_plan.md, in the BSF format.

The .docx is committed, so you only need this if you change the Markdown. The guidance in
<!-- ... --> becomes Word comments on the next paragraph, so it never counts as text.

    uv run --with "python-docx>=1.2" python make_docx.py
"""

from __future__ import annotations

import pathlib
import re

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Cm, Pt

HERE = pathlib.Path(__file__).resolve().parent


def add_runs(par, text: str) -> list:
    """**bold** spans become bold runs; everything else is plain."""
    runs = []
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if chunk:
            r = par.add_run(chunk)
            r.bold = i % 2 == 1
            runs.append(r)
    return runs


def build(src: pathlib.Path, out: pathlib.Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)                   # A4
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Cm(2.0))                                        # BSF §3: >= 2 cm
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Table Grid"):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(12) if name != "Heading 1" else Pt(14)          # BSF §3: 12 pt
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)

    pending: list[str] = []
    lines = src.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line:
            continue
        m = re.fullmatch(r"<!--\s*(.*?)\s*-->", line)
        if m:
            pending.append(m.group(1))
            continue
        if line.startswith("|"):
            rows = [line]
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i].rstrip())
                i += 1
            cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows
                     if not re.fullmatch(r"\|[-|\s]+\|", r)]
            table = doc.add_table(rows=len(cells), cols=len(cells[0]))
            table.style = "Table Grid"
            for r, row in enumerate(cells):
                for c, text in enumerate(row):
                    add_runs(table.cell(r, c).paragraphs[0], text)
            anchor = table.cell(0, 0).paragraphs[0].runs
        else:
            h = re.match(r"(#{1,3})\s+(.*)", line)
            if h:
                par = doc.add_heading(level=len(h.group(1)))
                anchor = add_runs(par, h.group(2))
            else:
                par = doc.add_paragraph()
                text = line
                # Consecutive non-blank lines are one paragraph, with a line break between.
                while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", "<!--")):
                    text += "\n" + lines[i].rstrip()
                    i += 1
                anchor = []
                for k, part in enumerate(text.split("\n")):
                    if k:
                        par.add_run().add_break()
                    anchor += add_runs(par, part)
        if pending and anchor:
            doc.add_comment(anchor, text="\n".join(pending), author="BSF guide", initials="BSF")
            pending = []
    doc.save(out)


if __name__ == "__main__":
    build(HERE / "research_plan.md", HERE / "research_plan.docx")
    print("wrote", HERE / "research_plan.docx")
